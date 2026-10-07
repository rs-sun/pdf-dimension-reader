"""R5 export/dims-override consistency diagnostics.

The helper is dump-only: it parses already-generated xlsx bytes and compares
them with the dims override rows that produced the workbook.
"""

from __future__ import annotations

import io
from typing import Any

import openpyxl

from export_gold_xlsx import (
    _collapse_gdt_rows,
    _detail_tol,
    _export_no,
    _first_present,
    _format_nominal,
)


SCHEMA_VERSION = "export_dim_override_consistency_v1"
CORE_FIELDS = (
    "no", "key", "nominal", "upper_tol", "lower_tol",
    "type", "page", "prefix", "unit",
)
BBOX_TOLERANCE_PT = 0.5


def build_export_dim_override_consistency_v1(
    *,
    dims_override: list[dict[str, Any]] | None,
    xlsx_bytes: bytes,
    pipeline_dims: list[dict[str, Any]] | None = None,
    pipeline_rerun_count: int = 0,
) -> dict[str, Any]:
    """Return row-level consistency diagnostics for one export."""
    dims = _collapse_gdt_rows(list(dims_override or []))
    pipeline_rows = _pipeline_rows_by_no(pipeline_dims or [])
    xlsx_rows = _parse_xlsx_rows(xlsx_bytes)

    rows = []
    matching_fields = 0
    required_fields = 0
    matching_no_positions = 0
    max_bbox_diff_pt = 0.0

    for idx, dim in enumerate(dims):
        override_row = _normalize_dim_row(dim, idx + 1)
        xlsx_row = xlsx_rows[idx] if idx < len(xlsx_rows) else None
        pipeline_row = pipeline_rows.get(str(override_row["no"]))
        diffs = []

        if xlsx_row and str(xlsx_row.get("no")) == str(override_row["no"]):
            matching_no_positions += 1

        for field in CORE_FIELDS:
            required_fields += 1
            expected = override_row.get(field)
            actual = (xlsx_row or {}).get(field)
            if _same_value(expected, actual):
                matching_fields += 1
            else:
                diffs.append({
                    "field": field,
                    "expected": expected,
                    "actual": actual,
                    "source": "xlsx",
                })

        bbox_diff = _bbox_max_abs_diff(override_row.get("bbox"), (pipeline_row or {}).get("bbox"))
        if bbox_diff is not None:
            max_bbox_diff_pt = max(max_bbox_diff_pt, bbox_diff)
            if bbox_diff > BBOX_TOLERANCE_PT:
                diffs.append({
                    "field": "bbox",
                    "expected": override_row.get("bbox"),
                    "actual": pipeline_row.get("bbox"),
                    "max_abs_diff_pt": round(bbox_diff, 6),
                    "source": "pipeline_dim_row",
                })

        rows.append({
            "stamp_id": override_row.get("stamp_id"),
            "dim_override_row": override_row,
            "pipeline_dim_row": pipeline_row,
            "xlsx_row": xlsx_row,
            "field_diffs": diffs,
            "used_dims_override": dims_override is not None,
            "pipeline_rerun_count": int(pipeline_rerun_count),
        })

    export_success = len(xlsx_bytes or b"") > 0 and len(xlsx_rows) >= len(dims)
    return {
        "schema_version": SCHEMA_VERSION,
        "used_dims_override": dims_override is not None,
        "pipeline_rerun_count": int(pipeline_rerun_count),
        "attempted_exports": 1,
        "successful_exports": 1 if export_success else 0,
        "export_success_rate": 1.0 if export_success else 0.0,
        "row_count": len(rows),
        "xlsx_row_count": len(xlsx_rows),
        "required_export_fields": required_fields,
        "matching_export_fields": matching_fields,
        "field_consistency_rate": (
            1.0 if required_fields == 0 else round(matching_fields / required_fields, 6)
        ),
        "matching_NO_positions": matching_no_positions,
        "stamp_count": len(dims),
        "no_order_consistency_rate": (
            1.0 if not dims else round(matching_no_positions / len(dims), 6)
        ),
        "max_bbox_diff_pt": round(max_bbox_diff_pt, 6),
        "bbox_tolerance_pt": BBOX_TOLERANCE_PT,
        "rows": rows,
    }


def summarize_export_dim_override_consistency(dump: dict[str, Any]) -> dict[str, Any]:
    """Return a compact HTTP-header-safe summary."""
    return {
        "schema_version": dump.get("schema_version", SCHEMA_VERSION),
        "used_dims_override": bool(dump.get("used_dims_override")),
        "pipeline_rerun_count": int(dump.get("pipeline_rerun_count") or 0),
        "field_consistency_rate": dump.get("field_consistency_rate"),
        "no_order_consistency_rate": dump.get("no_order_consistency_rate"),
        "export_success_rate": dump.get("export_success_rate"),
        "row_count": dump.get("row_count"),
        "diff_count": sum(len(row.get("field_diffs") or []) for row in dump.get("rows") or []),
    }


def _normalize_dim_row(dim: dict[str, Any], fallback_no: int) -> dict[str, Any]:
    return {
        "no": _export_no(dim, fallback_no),
        "key": "Y" if dim.get("is_key") else None,
        "nominal": _format_nominal(dim),
        "upper_tol": _detail_tol(dim.get("upper_tol")),
        "lower_tol": _detail_tol(dim.get("lower_tol")),
        "type": _first_present(dim, ["dim_type", "type"], ""),
        "page": _first_present(dim, ["page", "page_number", "pageIndex"], ""),
        "prefix": _first_present(dim, ["prefix"], ""),
        "unit": _first_present(dim, ["unit"], ""),
        "bbox": _normalize_bbox(dim.get("bbox")),
        "stamp_id": _first_present(dim, ["uuid", "stamp_uuid", "stamp_id", "id"], ""),
    }


def _parse_xlsx_rows(xlsx_bytes: bytes) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    report = wb["尺寸表"]
    detail = wb["数据明细"] if "数据明细" in wb.sheetnames else None
    rows = []
    for row_idx in range(3, report.max_row + 1):
        no = report.cell(row_idx, 1).value
        if no is None:
            continue
        detail_row = _detail_row_by_position(detail, len(rows) + 2)
        rows.append({
            "no": str(no),
            "key": report.cell(row_idx, 2).value,
            "nominal": report.cell(row_idx, 3).value,
            "upper_tol": detail_row.get("upper_tol"),
            "lower_tol": detail_row.get("lower_tol"),
            "type": detail_row.get("type", ""),
            "page": detail_row.get("page", ""),
            "prefix": detail_row.get("prefix", ""),
            "unit": detail_row.get("unit", ""),
            "bbox": None,
        })
    return rows


def _detail_row_by_position(detail, row_idx: int) -> dict[str, Any]:
    if detail is None or row_idx > detail.max_row:
        return {}
    return {
        "upper_tol": detail.cell(row_idx, 4).value,
        "lower_tol": detail.cell(row_idx, 5).value,
        "page": detail.cell(row_idx, 20).value or "",
        "type": detail.cell(row_idx, 21).value or "",
        "prefix": detail.cell(row_idx, 22).value or "",
        "unit": detail.cell(row_idx, 25).value or "",
    }


def _pipeline_rows_by_no(dims: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    rows = {}
    for idx, dim in enumerate(_collapse_gdt_rows(list(dims or []))):
        row = _normalize_dim_row(dim, idx + 1)
        rows[str(row["no"])] = row
    return rows


def _same_value(expected: Any, actual: Any) -> bool:
    if expected in (None, "") and actual in (None, ""):
        return True
    if isinstance(expected, (int, float)) or isinstance(actual, (int, float)):
        try:
            return abs(float(expected) - float(actual)) <= 1e-9
        except (TypeError, ValueError):
            return False
    return str(expected) == str(actual)


def _normalize_bbox(value: Any) -> dict[str, float] | None:
    if not isinstance(value, dict):
        return None
    try:
        return {
            "x": float(value.get("x", 0.0) or 0.0),
            "y": float(value.get("y", 0.0) or 0.0),
            "w": float(value.get("w", 0.0) or 0.0),
            "h": float(value.get("h", 0.0) or 0.0),
        }
    except (TypeError, ValueError):
        return None


def _bbox_max_abs_diff(
    expected: dict[str, float] | None,
    actual: dict[str, float] | None,
) -> float | None:
    if expected is None or actual is None:
        return None
    return max(abs(float(expected[key]) - float(actual[key])) for key in ("x", "y", "w", "h"))
