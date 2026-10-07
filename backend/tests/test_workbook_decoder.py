"""Synthetic workbook checks, with no drawings or saved binary fixtures.

Run from the candidate root with:
    PYTHONDONTWRITEBYTECODE=1 python -B -m pytest -p no:cacheprovider backend/tests

The XLS adapter contract uses in-memory doubles; the malformed-XLS check uses
the installed decoder. A separate private audit compared real synthetic BIFF
bytes with the previous decoder. Its temporary writer is not a test dependency.
"""

import builtins
import io
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest


sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from _inspection_workbook_decoder import (  # noqa: E402
    ParsedSheet,
    decode_workbook_sheets,
    validate_sheet_bounds,
)


def _xlsx_bytes():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = '合成表'
    worksheet['D3'] = '长度'
    worksheet['D4'] = 12.5
    worksheet['E4'] = True
    worksheet['F4'] = datetime(2026, 1, 2, 15, 30)
    worksheet['G4'] = '=1+2'
    workbook.create_sheet('空表')
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('extension', ['.xlsx', '.xlsm'])
def test_zip_workbooks_keep_openpyxl_coordinates_and_data_only(extension):
    sheets = decode_workbook_sheets('synthetic' + extension, _xlsx_bytes())
    assert [sheet.name for sheet in sheets] == ['合成表', '空表']
    sheet = sheets[0]
    assert (sheet.nrows, sheet.ncols) == (4, 7)
    assert sheet.cell(2, 3) == '长度'
    assert sheet.cell(3, 3) == 12.5
    assert sheet.cell(3, 4) is True
    assert sheet.cell(3, 5) == datetime(2026, 1, 2, 15, 30)
    assert sheet.cell(3, 6) is None  # no fabricated formula result
    assert sheet.cell(0, 0) is None
    assert sheet.cell(20, 20) is None
    assert (sheets[1].nrows, sheets[1].ncols) == (0, 0)


def _fake_calamine(monkeypatch, matrices, error=None):
    class CalamineError(Exception):
        pass

    class Sheet:
        def __init__(self, rows):
            self.rows = rows

        def to_python(self, *, skip_empty_area):
            assert skip_empty_area is False
            return [list(row) for row in self.rows]

    class Workbook:
        sheet_names = list(matrices)
        closed = False

        @classmethod
        def from_filelike(cls, stream):
            assert isinstance(stream, io.BytesIO)
            if error:
                raise CalamineError(error)
            return cls()

        def __enter__(self):
            return self

        def __exit__(self, *args):
            Workbook.closed = True

        def get_sheet_by_name(self, name):
            return Sheet(matrices[name])

    monkeypatch.setitem(sys.modules, 'python_calamine', SimpleNamespace(
        CalamineError=CalamineError, CalamineWorkbook=Workbook,
    ))
    return Workbook


def test_xls_adapter_preserves_absolute_coordinates_and_independent_sheets(monkeypatch):
    workbook = _fake_calamine(monkeypatch, {
        '第一表': [['', '', ''], ['', '', '文字'], ['', '', 2.25], ['', '', True]],
        '第二表': [[False, '其他']],
        '空表': [],
    })
    sheets = decode_workbook_sheets('synthetic.xls', b'synthetic adapter input')
    assert [(s.name, s.nrows, s.ncols) for s in sheets] == [
        ('第一表', 4, 3), ('第二表', 1, 2), ('空表', 0, 0),
    ]
    assert sheets[0].cell(1, 2) == '文字'
    assert sheets[0].cell(2, 2) == 2.25
    assert type(sheets[0].cell(3, 2)) is int
    assert sheets[0].cell(3, 2) == 1
    assert sheets[1].cell(0, 0) == 0
    assert sheets[0].cell(0, 0) == ''
    assert sheets[0].cell(-1, 0) is None
    assert sheets[2].cell(0, 0) is None
    assert workbook.closed


@pytest.mark.parametrize('value', [
    date(2026, 1, 2), datetime(2026, 1, 2, 15, 30), time(6), timedelta(hours=2),
])
def test_xls_date_time_and_duration_fail_closed(monkeypatch, value):
    workbook = _fake_calamine(monkeypatch, {'合成表': [[12.5, value]]})
    with pytest.raises(ValueError, match='另存为 .xlsx'):
        decode_workbook_sheets('synthetic.xls', b'synthetic adapter input')
    assert workbook.closed


def test_calamine_parser_failure_is_a_client_workbook_error(monkeypatch):
    _fake_calamine(monkeypatch, {}, error='synthetic invalid workbook')
    with pytest.raises(ValueError, match='无法解析测量结果工作簿'):
        decode_workbook_sheets('synthetic.xls', b'invalid')


def test_missing_xls_dependency_remains_a_server_configuration_error(monkeypatch):
    real_import = builtins.__import__

    def import_without_calamine(name, *args, **kwargs):
        if name == 'python_calamine':
            raise ImportError('synthetic missing dependency')
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, '__import__', import_without_calamine)
    with pytest.raises(RuntimeError, match='python-calamine==0.8.2'):
        decode_workbook_sheets('synthetic.xls', b'invalid')


def test_real_calamine_rejects_malformed_xls():
    pytest.importorskip('python_calamine')
    with pytest.raises(ValueError, match='无法解析测量结果工作簿'):
        decode_workbook_sheets('synthetic.xls', b'not an Excel workbook')


def test_xlsx_cannot_bypass_its_decoder_by_renaming_to_xls():
    pytest.importorskip('python_calamine')
    with pytest.raises(ValueError, match='正确扩展名'):
        decode_workbook_sheets('synthetic.xls', _xlsx_bytes())


@pytest.mark.parametrize('extension', ['.xlsx', '.xlsm'])
def test_corrupt_zip_workbooks_raise_value_error(extension):
    with pytest.raises(ValueError, match='无法解析测量结果工作簿'):
        decode_workbook_sheets('synthetic' + extension, b'not a ZIP workbook')


@pytest.mark.parametrize('extension', ['.csv', '.ods', ''])
def test_unsupported_extensions_are_rejected(extension):
    with pytest.raises(ValueError, match='不支持的测量文件类型'):
        decode_workbook_sheets('synthetic' + extension, b'')


@pytest.mark.parametrize('sheets, message', [
    ([ParsedSheet('表', 2, 2, None)] * 3, 'sheet 数量过多'),
    ([ParsedSheet('表', 4, 2, None)], '行数过多'),
    ([ParsedSheet('表', 2, 4, None)], '列数过多'),
])
def test_workbook_resource_limits(sheets, message):
    with pytest.raises(ValueError, match=message):
        validate_sheet_bounds(sheets, 'synthetic.xlsx',
                              max_workbook_sheets=2, max_sheet_rows=3, max_sheet_cols=3)


def test_workbook_resource_limit_boundary_is_allowed():
    validate_sheet_bounds([ParsedSheet('表', 3, 3, None)] * 2, 'synthetic.xlsx',
                          max_workbook_sheets=2, max_sheet_rows=3, max_sheet_cols=3)
